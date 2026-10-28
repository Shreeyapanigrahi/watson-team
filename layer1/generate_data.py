"""Synthetic corpus generator (Section 9.5).

Builds real PDF / DOCX / PPTX / TXT files from seed vocabularies and text
templates, so the parser has genuine files to process. Nothing here is a hand-typed
data dump: each document is parameters (brand, region, year, sample size, ...)
filled into templates by a seeded random generator.

Planted cases the pipeline must handle:
  - 1 scanned (image-only) PDF           -> parse warning
  - 1 KOL file with personal data        -> PII detection
  - 1 near-duplicate pair                -> hold for review
  - 1 new-version pair (same title/author, later date)
  - 1 exact duplicate                    -> rejected
  - 2 untitled files, 1 non-UTF-8 text file, 2 image-only slides

Sentiment is designed with variance (Section 9.4): two brands fall from >0.6 to
<0.4 after 2022, two stay high, two stay low, several flip by geography, one has
a single geography.

Run:  python -m layer1.generate_data [--out DIR]
Ground truth is written to manifest.json. The ingestion pipeline never reads it;
only the seed loader (bulk_ingest.py) and the tests do.
"""
import argparse
import copy
import io
import json
import random
import shutil
from datetime import date, timedelta
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from docx.shared import Pt
from pptx import Presentation
from pptx.util import Inches, Pt as PPt
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from . import config, vocab

SEED = 42
FOOTER = "Synthetic document generated for the Watson prototype. All brand names and figures are fictional."
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]

PROFILE = {
    "Opdivra": "trajectory", "Cosentex": "trajectory",
    "Tagrisa": "high", "Dupixa": "high",
    "Trodelvia": "low", "Verquvia": "low",
    "Keytrava": "us_pos", "Jardiva": "us_pos",
    "Enhertra": "eu_pos", "Humivex": "eu_pos",
    "Leqembra": "us_only",
}


def sentiment(brand, region, year, rng):
    p = PROFILE.get(brand, "mid")
    if p == "trajectory":
        base = 0.72 if year <= 2022 else 0.31
    elif p == "high":
        base = 0.75
    elif p == "low":
        base = 0.30
    elif p == "us_pos":
        base = 0.72 if region == "US" else 0.34
    elif p == "eu_pos":
        base = 0.72 if region == "EU" else 0.34
    elif p == "us_only":
        base = 0.62
    else:
        base = 0.55
    return round(min(0.95, max(0.05, base + rng.uniform(-0.05, 0.05))), 2)


def sent_adj(s):
    return ("strongly positively" if s >= 0.65 else "moderately positively" if s >= 0.5
            else "neutrally" if s >= 0.4 else "negatively")


def specialist(ta, ind):
    if ta == "Immunology":
        return "dermatologist" if ind in ("plaque psoriasis", "atopic dermatitis") else "rheumatologist"
    return {"Oncology": "oncologist", "Cardiology": "cardiologist", "Neurology": "neurologist"}[ta]


def countries_of(region):
    return {"EU": ["DE", "FR", "IT", "ES", "UK"], "US": ["US"], "APAC": ["JP"]}[region]


def cname(code):
    return vocab.GEOGRAPHIES[code]["name"]


BARRIERS = ["cost and formulary access remain barriers", "infusion logistics add to treatment burden",
            "long-term safety data are still limited", "the dosing schedule complicates adherence",
            "payer prior-authorisation requirements slow initiation"]
ADH = ["missed maintenance doses", "treatment burden over long durations", "injection-site discomfort",
       "monitoring requirements", "discontinuation after early adverse events"]


class Gen:
    def __init__(self, rng):
        self.rng = rng

    def pick(self, seq):
        return self.rng.choice(seq)

    def ctx(self, spec, brand=None):
        r = self.rng
        brand = brand or spec["brands"][0]
        b = vocab.BRANDS[brand]
        ta = b["ta"]
        ind = spec.get("indication") or b["indications"][0]
        code = spec.get("country") or self.pick(countries_of(spec["region"]) if spec["region"] != "GLOBAL" else ["US", "DE"])
        ae = r.sample(vocab.ADVERSE_EVENTS[ta], 2)
        s = sentiment(brand, spec["region"] if spec["region"] != "GLOBAL" else "US", spec["year"], r)
        return {
            "brand": brand, "drug": b["drug"], "dclass": b["drug_class"], "mech": b["mechanism"].lower(),
            "ind": ind, "ta": ta, "country": cname(code), "spec": specialist(ta, ind),
            "spec_cap": specialist(ta, ind).capitalize(), "ae1": ae[0], "ae2": ae[1],
            "adj": sent_adj(s), "score": s, "year": spec["year"], "n": spec.get("n", 0),
            "month": MONTHS[r.randint(0, 5)], "month2": MONTHS[r.randint(6, 11)],
            "pct1": r.randint(48, 86), "pct2": r.randint(18, 52), "pct3": r.randint(35, 78),
            "intent": self.pick(["high", "moderate", "cautious"]), "line": self.pick(["first-line", "second-line", "refractory disease"]),
            "barrier": self.pick(BARRIERS), "adh": self.pick(ADH), "k": r.randint(4, 25),
            "mins": r.randint(18, 35), "company": b["company"], "launch": r.randint(2017, 2022),
        }


def f(t, c):
    return t.format(**c)


# ------------------------------------------------------------------ study-type builders
def build_hcp(g, spec):
    c = g.ctx(spec)
    c["countries"] = ", ".join(cname(x) for x in countries_of(spec["region"]))
    rich = spec["richness"]
    sec = []
    sec.append(("Executive Summary", [
        f(g.pick(["This survey of {n} {spec}s across {countries} examined how clinicians perceive {brand} ({drug}) for {ind}. Fieldwork took place in {month} {year}.",
                  "Between {month} and {month2} {year}, {n} {spec}s in {countries} completed a perception survey on {brand} ({drug}), a {dclass}, in {ind}."]), c),
        f("Overall HCP sentiment index for {brand} in {country}: {score:.2f} (scale 0 to 1, assessed {year}). Respondents viewed the brand {adj}.", c),
    ]))
    if rich in ("med", "high"):
        sec.append(("Methodology", [
            f("Respondents were {spec}s who treat at least {k} patients per month with {ind}. The sample size was n={n}. Interviews lasted {mins} minutes and were conducted online.", c),
            "Results are weighted to national prescribing volumes. Statistical significance was assessed at the 95% confidence level.",
        ]))
    kf = []
    pool_clin = [
        "{spec_cap}s in {country} rated {brand} tolerability profile as {adj}, citing {ae1} and {ae2} as the main concerns among patients with {ind}.",
        "{pct1}% of respondents considered the efficacy of {brand} clinically meaningful versus the current standard of care, with the strongest support in {line} use.",
        "Prescribing intent for {brand} among {spec}s was {intent}, with {pct2}% planning to increase use in the next 12 months.",
        "Respondents highlighted the mechanism of action of {brand} ({mech}) as a differentiator in {ind}, although {barrier}.",
        "Adherence concerns centred on {adh}, particularly in patients with comorbidities and polypharmacy.",
        "Real-world evidence and patient-reported outcomes were requested by {pct3}% of respondents before expanding use beyond first-line therapy.",
    ]
    pool_plain = [
        "Respondents in {country} had mixed views on {brand}.", "Many {spec}s said they would like more information about {brand}.",
        "Opinions about {brand} varied between respondents.", "Several respondents commented on the product overall.",
    ]
    if rich == "low":
        kf = [f(t, c) for t in g.rng.sample(pool_plain, 3)]
    elif rich == "med":
        kf = [f(t, c) for t in g.rng.sample(pool_clin, 3)] + [f(g.pick(pool_plain), c)]
    else:
        kf = [f(t, c) for t in g.rng.sample(pool_clin, 5)]
    sec.append(("Key Findings", kf))
    table = None
    if rich == "high":
        sec.append(("Tolerability and Safety Perceptions", [
            f("Tolerability was assessed against {ae1} and {ae2}. {spec_cap}s reported managing these events with dose adjustment and supportive care.", c),
            f("Safety perception differed by experience: respondents with more than {k} treated patients rated {brand} more favourably.", c),
        ]))
        table = [["Segment", "Respondents", "Sentiment index"],
                 ["High-volume prescribers", str(int(spec["n"] * 0.4)), f"{min(0.95, c['score'] + 0.04):.2f}"],
                 ["Low-volume prescribers", str(int(spec["n"] * 0.6)), f"{max(0.05, c['score'] - 0.03):.2f}"]]
    if rich != "low":
        sec.append(("Conclusions", [f("{brand} is perceived {adj} in {country}. Priorities for the brand team are {adh} and evidence generation for {ind}.", c)]))
    spec["sentiment"] = [{"brand": c["brand"], "region": spec["region"], "year": spec["year"], "score": c["score"]}]
    spec["indications"] = [c["ind"]]
    return sec, table, [f("Speaker note: sentiment for {brand} was {score:.2f}.", c)]


def build_cl(g, spec):
    """Competitive landscape: every brand in the indication, with share and sentiment."""
    ind, region, year = spec["indication"], spec["region"], spec["year"]
    ta = vocab.INDICATIONS[ind]["ta"]
    brands = [b for b, v in vocab.BRANDS.items() if ind in v["indications"]]
    spec["brands"] = brands
    sec, sents, shares = [], [], g.rng.sample(range(8, 40), len(brands))
    rc = region if region != "GLOBAL" else "US"
    sec.append(("Market Overview", [
        f"In the {region} {ind} market, {len(brands)} brands compete for {g.pick(['second-line', 'first-line'])}+ patients as of Q{g.rng.randint(1, 4)} {year}. "
        f"The segment is shaped by guideline updates, reimbursement decisions and real-world evidence.",
        f"The number of branded options in {ind} has grown since {year - g.rng.randint(2, 4)}, increasing pressure on standard of care positioning."]))
    for b, sh in zip(brands, shares):
        c = g.ctx({**spec, "region": rc}, brand=b)
        sents.append({"brand": b, "region": region, "year": year, "score": c["score"]})
        sec.append((f"Positioning: {b}", [
            f"In the {region} {ind} market, {b} maintains a {sh}% market share as of Q{g.rng.randint(1, 4)} {year}, with the strongest uptake among {c['spec']}s.",
            f"Sentiment index for {b} in {region}: {c['score']:.2f} (scale 0 to 1, {year}). Clinicians view the brand {c['adj']}; {c['barrier']}.",
        ]))
    lead, chal = brands[0], brands[-1]
    sec.append(("Competitive Dynamics", [
        f"{chal} has been gaining ground since its {g.rng.randint(2020, 2023)} approval for {ind}, pressuring {lead} in {g.pick(['first-line', 'second-line'])} use.",
        "Price competition, biosimilar entry and head-to-head evidence are expected to reshape formulary position over the next 24 months."]))
    sec.append(("Outlook", [f"We expect continued consolidation in {ind}. Differentiation will depend on tolerability, dosing convenience and evidence of overall survival or clinical benefit."]))
    spec["sentiment"], spec["indications"] = sents, [ind]
    spec["therapeutic_area"] = ta
    return sec, None, [f"Speaker note: market shares for {len(brands)} brands in {region}."]


def build_id(g, spec):
    ind, year = spec["indication"], spec["year"]
    ta = vocab.INDICATIONS[ind]["ta"]
    brands = [b for b, v in vocab.BRANDS.items() if ind in v["indications"]]
    spec["brands"] = brands
    c = g.ctx(spec, brand=brands[0])
    sec = [
        ("Disease Overview", [f"{ind.capitalize()} represents {g.rng.randint(10, 25)}-{g.rng.randint(26, 40)}% of cases in its category. Unmet need remains significant in patients who have progressed on first-line therapy, with an estimated overall survival of approximately {g.rng.randint(20, 60)}% at 36 months.",
                              f"Prevalence and incidence vary by region; {spec['region']} data show {g.pick(['rising', 'stable'])} diagnosis rates since {year - 2}."]),
        ("Standard of Care", [f"Current guidelines recommend {g.pick(['combination therapy', 'targeted therapy', 'monotherapy'])} as first-line treatment. Comorbidity management and biomarker testing influence the treatment choice.",
                              f"Marketed options include {', '.join(brands)}."]),
        ("Unmet Need", [f"{c['spec_cap']}s identified {c['adh']} and {c['barrier']} as the main unmet needs in {ind}.",
                        "Patient-reported outcomes and quality of life measures are increasingly used to guide treatment decisions."]),
        ("Pipeline and Emerging Therapies", [f"Several phase 3 programmes are evaluating next-generation options in {ind}. Early data suggest improved progression-free survival and a manageable safety profile."]),
        ("Implications for Brands", [f"{brands[0]} is positioned to benefit if payers expand reimbursement. Brand teams should invest in real-world evidence and {c['spec']} education."]),
    ]
    spec["sentiment"], spec["indications"], spec["therapeutic_area"] = [], [ind], ta
    return sec, None, []


def build_ma(g, spec):
    brand = spec["brands"][0]
    c = g.ctx(spec, brand=brand)
    c["body"] = vocab.GEOGRAPHIES[spec["country"]]["regulatory_body"].split(" / ")[0]
    c["payer"] = vocab.GEOGRAPHIES[spec["country"]]["payer_type"].lower()
    c["months"] = g.pick([12, 18, 24, 36])
    c["until"] = spec["year"] + 2
    c["risk"] = g.pick(["low", "moderate", "high"])
    sec = [
        ("Access Overview", [f("Reimbursement for {brand} in {country} was approved on a managed entry agreement requiring real-world evidence submission within {months} months of launch under {body}.", c),
                             f("The {payer} system applies health technology assessment before listing {brand} for {ind}.", c)]),
        ("Payer Perspectives", [f("Payers in {country} view the price of {brand} as acceptable only with a demonstrated clinical benefit versus the standard of care.", c),
                                "Formulary placement is reviewed annually, and pricing is linked to outcomes in the registry."]),
        ("Evidence Requirements", [f("Required evidence includes overall survival or response rate data, patient-reported outcomes and comparative effectiveness against {brand}'s main competitor.", c)]),
        ("Pricing and Reimbursement Outlook", [f("A confidential rebate is expected to remain in place until {until}. Market access risk is {risk} for new indications.", c)]),
    ]
    spec["sentiment"], spec["indications"], spec["therapeutic_area"] = [], [c["ind"]], c["ta"]
    spec["indication"] = c["ind"]
    return sec, None, []


def build_kol(g, spec):
    brand = spec["brands"][0]
    c = g.ctx(spec, brand=brand)
    kol = f"Dr. {g.pick(vocab.PERSON_FIRST)} {g.pick(vocab.PERSON_LAST)}"
    if spec.get("encoding"):
        kol = "Dr. Jürgen Müller"          # non-ASCII name so the non-UTF-8 file really needs its encoding detected
    c["kol"] = kol
    intro = f"{kol}, a leading {c['spec']} in {c['country']}, was interviewed on {MONTHS[g.rng.randint(0, 11)]} {spec['year']} about {brand} ({c['drug']}) in {c['ind']}."
    if spec.get("pii"):
        intro = ("Dr. John Martinez (DOB: 14/03/1968, email: j.martinez@clinic.de, phone: +49 30 5551 2847) was interviewed on "
                 f"{MONTHS[g.rng.randint(0, 11)]} {spec['year']} about {brand} ({c['drug']}) in {c['ind']}.")
        c["kol"] = "Dr. John Martinez"
    sec = [
        ("Interview Summary", [intro,
                               f("{kol} noted that while the efficacy data for {brand} are compelling, {barrier}.", {**c, "barrier": c["barrier"]})]),
        ("Clinical Perspective", [f("{kol} said that tolerability, particularly {ae1} and {ae2}, shapes the choice between {brand} and alternatives in {ind}.", c),
                                  f("In {country}, {kol} explained that treatment decisions follow national guidelines and local payer rules.", c)]),
        ("Advice to the Brand Team", [f("{kol} recommended investing in real-world evidence and in education on {adh}.", c)]),
    ]
    if spec.get("pii"):
        sec[0] = (sec[0][0], [intro, f"Dr. John Martinez noted that while the efficacy data for {brand} are compelling, {c['barrier']}."])
        sec[1] = (sec[1][0], [f"Dr. Martinez said that tolerability, particularly {c['ae1']} and {c['ae2']}, shapes the choice between {brand} and alternatives in {c['ind']}.",
                              f"In {c['country']}, treatment decisions follow national guidelines and local payer rules."])
        sec[2] = (sec[2][0], [f"Dr. Martinez recommended investing in real-world evidence and in education on {c['adh']}."])
    spec["sentiment"], spec["indications"], spec["therapeutic_area"] = [], [c["ind"]], c["ta"]
    spec["indication"] = c["ind"]
    return sec, None, []


def build_tal(g, spec):
    ta, region, year = spec["ta"], spec["region"], spec["year"]
    pathway = {"Oncology": "PD-1 and KRAS G12C pathways", "Cardiology": "sGC and SGLT2 pathways",
               "Neurology": "CGRP and anti-amyloid pathways", "Immunology": "IL-17 and IL-23 pathways"}[ta]
    brands = [b for b, v in vocab.BRANDS.items() if v["ta"] == ta]
    spec["brands"] = brands
    inds = sorted({i for b in brands for i in vocab.BRANDS[b]["indications"]})
    sec = [
        ("Pipeline Overview", [f"The {ta.lower()} pipeline shows {g.rng.randint(8, 16)} late-stage assets targeting {pathway}, creating significant competitive pressure on established biologics."]),
        ("Marketed Landscape", [f"Marketed brands include {', '.join(brands)}, spanning {', '.join(inds[:4])}."]),
        ("Key Takeaways", [f"First-line positioning is contested in {region}. Payer scrutiny and biosimilar competition continue to rise.",
                           "Mechanism of action differentiation and dosing convenience will drive prescribing."]),
        ("Implications", [f"Brand teams in {ta} should prioritise evidence for {inds[0]} and monitor phase 3 readouts through {year + 1}."]),
    ]
    spec["sentiment"], spec["indications"], spec["therapeutic_area"] = [], inds, ta
    return sec, None, []


EXTRA = [
    "Clinical guidelines in {country} recommend {line} use of agents in this class, and recent updates have broadened eligibility based on biomarker status and comorbidity profile.",
    "Real-world evidence from registries suggests that outcomes in routine practice are broadly consistent with phase 3 trial results, although persistence and discontinuation rates are higher outside clinical trials.",
    "{spec_cap}s emphasised that patient-reported outcomes and quality of life are increasingly important when choosing between {dclass} options.",
    "Treatment burden, including monitoring visits and {adh}, was described as a key factor in long-term adherence and persistence.",
    "Combination therapy and sequencing strategies remain debated; respondents noted limited head-to-head evidence between {brand} and competing agents.",
    "Safety monitoring focused on {ae1} and {ae2}, with dose interruption or discontinuation used for severe adverse events.",
    "Access to specialist centres influences adoption: academic centres adopt new therapies faster than community practices, where formulary restrictions and reimbursement criteria apply.",
    "Education needs include biomarker testing, management of adverse events and counselling patients on the expected benefit-risk.",
    "Payer decisions increasingly rely on health technology assessment, requiring evidence of clinical benefit versus the standard of care and cost-effectiveness analyses.",
    "The unmet need in {ind} persists in patients who relapse or become refractory, where options are limited and outcomes remain poor.",
    "Subgroup analyses suggest that benefit with {brand} is most pronounced in patients with a favourable biomarker profile and a lower comorbidity burden.",
    "Sequencing after progression remains an open question; {spec}s in {country} reported using a different mechanism of action after {brand} fails.",
    "Dose modification guidance for {brand} is generally followed, but {spec}s said that practical monitoring of {ae1} depends on local capacity and patient travel distance.",
    "Cost-effectiveness submissions typically model {line} use, and payers question the durability of benefit beyond the trial follow-up period.",
    "Peer-to-peer education and congress data were cited as the most trusted sources when {spec}s evaluate new evidence on {dclass} agents.",
    "Guideline committees are expected to review the evidence base for {ind} again next year, which could change eligibility criteria and referral patterns.",
    "Patient advocacy groups are pushing for earlier access, and shared decision-making tools are being piloted to explain benefit, risk and treatment burden.",
    "Supply and logistics are not a concern for {brand}, though cold-chain handling and nurse-led administration add workload in busy clinics.",
    "Biosimilar and generic competition in adjacent classes is shaping price expectations, even where {brand} has no direct substitute.",
    "Regional variation within {country} is notable: large centres report faster uptake and broader use than smaller regional hospitals.",
    "Respondents expect real-world persistence on {brand} to be lower than in trials, mainly because of {adh} and the complexity of managing {ae2}.",
    "Medical education priorities include recognising {ae1} early, communicating realistic expectations, and coordinating with primary care.",
    "Registry data show that outcomes in {ind} depend heavily on early diagnosis, access to specialist care and adherence to maintenance therapy.",
    "The competitive set is evolving as new mechanisms enter phase 3 development, with readouts expected over the next two to three years.",
]


def pad(g, spec, sections, fmt):
    """Make document length and clinical density vary with `richness`
    (low: none, med: +8 paragraphs per section, high: +13; slides capped at +3).
    Real market research reports run to many pages, so short documents would
    make the chunking comparison meaningless."""
    k = {"low": 0, "med": 8, "high": 13}[spec["richness"]]
    if fmt == "pptx":
        k = min(k, 3)
    out = []
    for i, (h, paras) in enumerate(sections):
        if k and i > 0 and not h.startswith("Positioning"):
            c = g.ctx(spec, brand=spec["brands"][0])
            paras = paras + [f(t, c) for t in g.rng.sample(EXTRA, k)]
        out.append((h, paras))
    return out


BUILDERS = {"HCP Perception Survey": build_hcp, "Competitive Landscape Report": build_cl,
            "Indication Deep Dive": build_id, "Market Access Insight": build_ma,
            "KOL Interview Summary": build_kol, "Therapy Area Landscape": build_tal}


# ------------------------------------------------------------------ renderers
def _meta_line(spec):
    return f"Report date: {spec['report_date']}" + (f" | Prepared by: {spec['author']}" if spec.get("author") else "")


def render_pdf(spec, sections, table, path):
    ss = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=spec["title"] if spec.get("meta_title", True) else "",
                            author=spec.get("author") or "")
    flow = []
    if spec.get("show_title", True):
        flow += [Paragraph(spec["title"], ss["Title"])]
    flow += [Paragraph(_meta_line(spec), ss["Normal"]), Spacer(1, 10)]
    for h, paras in sections:
        flow.append(Paragraph(h, ss["Heading2"]))
        for p in paras:
            flow += [Paragraph(p, ss["BodyText"]), Spacer(1, 4)]
    if table:
        t = Table(table)
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey)]))
        flow += [t, Spacer(1, 8)]
    flow.append(Paragraph(FOOTER, ss["Italic"]))
    doc.build(flow)


def render_scanned_pdf(spec, sections, table, path):
    tmp = Path(str(path) + ".tmp.pdf")
    render_pdf(spec, sections, table, tmp)
    src, out = pymupdf.open(tmp), pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=100)
        p = out.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=pix.tobytes("png"))
    out.save(path)
    out.close()
    src.close()
    tmp.unlink()


def render_docx(spec, sections, table, path):
    d = DocxDocument()
    if spec.get("show_title", True):
        d.add_heading(spec["title"], 0)
    cp = d.core_properties
    cp.title = spec["title"] if spec.get("meta_title", True) else ""
    cp.author = spec.get("author") or ""
    y, m, dd = map(int, spec["report_date"].split("-"))
    from datetime import datetime
    cp.created = datetime(y, m, dd, 9, 0, 0)
    d.add_paragraph(_meta_line(spec))
    for h, paras in sections:
        d.add_heading(h, 1)
        for p in paras:
            d.add_paragraph(p)
        if h == "Tolerability and Safety Perceptions" and table:
            t = d.add_table(rows=len(table), cols=len(table[0]))
            t.style = "Table Grid"
            for i, row in enumerate(table):
                for j, v in enumerate(row):
                    t.cell(i, j).text = v
    d.add_paragraph(FOOTER)
    d.save(path)


def _chart_png(label):
    img = Image.new("RGB", (640, 360), "white")
    dr = ImageDraw.Draw(img)
    for i, h in enumerate([120, 200, 160, 260, 90]):
        dr.rectangle([60 + i * 100, 320 - h, 120 + i * 100, 320], fill=(40, 110, 160))
    dr.text((20, 10), label, fill="black")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    return buf


def render_pptx(spec, sections, table, notes, path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = spec["title"]
    s.placeholders[1].text = _meta_line(spec)
    cp = prs.core_properties
    cp.title = spec["title"] if spec.get("meta_title", True) else ""
    cp.author = spec.get("author") or ""
    from datetime import datetime
    y, m, dd = map(int, spec["report_date"].split("-"))
    cp.created = datetime(y, m, dd, 9, 0, 0)
    for i, (h, paras) in enumerate(sections):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = h
        tf = s.placeholders[1].text_frame
        tf.text = paras[0]
        for p in paras[1:]:
            tf.add_paragraph().text = p
        for para in tf.paragraphs:
            for r in para.runs:
                r.font.size = PPt(16)
        if i == 0 and notes:
            s.notes_slide.notes_text_frame.text = " ".join(notes)
    if spec.get("image_slide"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.shapes.add_picture(_chart_png("Market share by brand"), Inches(2), Inches(1.5), width=Inches(8))
    if spec.get("unordered_boxes"):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        lower = s.shapes.add_textbox(Inches(1), Inches(4.5), Inches(10), Inches(1))
        lower.text_frame.text = "Second takeaway: payer scrutiny continues to rise."
        upper = s.shapes.add_textbox(Inches(1), Inches(1.5), Inches(10), Inches(1))
        upper.text_frame.text = "First takeaway: differentiation depends on dosing convenience."
    last = prs.slides.add_slide(prs.slide_layouts[1])
    last.shapes.title.text = "Disclaimer"
    last.placeholders[1].text_frame.text = FOOTER
    prs.save(path)


def render_txt(spec, sections, table, path):
    lines = []
    if spec.get("show_title", True):
        lines += [spec["title"], ""]
    lines += [_meta_line(spec), ""]
    for h, paras in sections:
        lines += [h.upper(), ""] + [p + "\n" for p in paras]
    lines.append(FOOTER)
    text = "\n".join(lines)
    enc = spec.get("encoding", "utf-8")
    Path(path).write_bytes(text.encode(enc, errors="replace"))


# ------------------------------------------------------------------ plan (parameters, not data)
# (study_type, key fields). ingest_age: days before "now" the seed loader back-dates ingestion.
HCP = [
    dict(brand="Opdivra", year=2021, region="EU", n=320, fmt="pdf", rich="high", age=700),
    dict(brand="Opdivra", year=2022, region="US", n=210, fmt="docx", rich="med", sens="RESTRICTED"),
    dict(brand="Opdivra", year=2023, region="EU", n=300, fmt="pdf", rich="high", age=745),
    dict(brand="Opdivra", year=2024, region="US", n=45, fmt="docx", rich="med"),
    dict(brand="Cosentex", year=2022, region="EU", n=350, fmt="pdf", rich="high"),
    dict(brand="Cosentex", year=2024, region="EU", n=120, fmt="pdf", rich="med", scanned=True),
    dict(brand="Keytrava", year=2023, region="US", n=500, fmt="docx", rich="high", sens="RESTRICTED", near_dup_source=True),
    dict(brand="Keytrava", year=2023, region="EU", n=280, fmt="pdf", rich="med", no_author=True, show_title=False, meta_title=False),
    dict(brand="Jardiva", year=2023, region="US", n=18, fmt="pdf", rich="low", no_author=True),
    dict(brand="Enhertra", year=2024, region="EU", n=260, fmt="docx", rich="high"),
    dict(brand="Trodelvia", year=2022, region="US", n=150, fmt="pdf", rich="med"),
    dict(brand="Entrestra", year=2021, region="EU", n=90, fmt="docx", rich="low", sens="RESTRICTED", exact_dup_source=True),
]
CL = [
    dict(ind="non-small cell lung cancer", region="EU", year=2023, fmt="pptx", new_version_source=True, image_slide=True),
    dict(ind="non-small cell lung cancer", region="US", year=2023, fmt="pdf"),
    dict(ind="non-small cell lung cancer", region="EU", year=2024, fmt="pdf", sens="CONFIDENTIAL"),
    dict(ind="heart failure with reduced ejection fraction", region="GLOBAL", year=2024, fmt="pptx", unordered_boxes=True),
    dict(ind="plaque psoriasis", region="EU", year=2022, fmt="pdf"),
    dict(ind="plaque psoriasis", region="US", year=2024, fmt="pptx"),
    dict(ind="HER2-positive breast cancer", region="EU", year=2023, fmt="pdf"),
    dict(ind="atopic dermatitis", region="US", year=2023, fmt="pptx"),
]
ID = [dict(ind=i, region=r, year=y, fmt="docx") for i, r, y in [
    ("EGFR-mutant NSCLC", "EU", 2023), ("HER2-positive breast cancer", "EU", 2024),
    ("heart failure with reduced ejection fraction", "US", 2022), ("atrial fibrillation", "EU", 2023),
    ("multiple sclerosis", "EU", 2024), ("migraine", "US", 2023), ("plaque psoriasis", "US", 2022)]]
MA = [dict(brand=b, country=c, year=y, fmt="pdf") for b, c, y in [
    ("Keytrava", "DE", 2023), ("Enhertra", "FR", 2024), ("Jardiva", "IT", 2022),
    ("Ocrevia", "ES", 2023), ("Skyrixa", "UK", 2024), ("Humivex", "US", 2022)]]
KOL = [
    dict(brand="Aimovia", country="DE", year=2024, fmt="docx", age=330),
    dict(brand="Eliquex", country="FR", year=2023, fmt="docx"),
    dict(brand="Tagrisa", country="DE", year=2024, fmt="docx", pii=True),
    dict(brand="Ocrevia", country="UK", year=2023, fmt="txt", age=400),
    dict(brand="Verquvia", country="US", year=2024, fmt="txt", show_title=False, no_author=True),
    dict(brand="Lumakro", country="ES", year=2024, fmt="txt", encoding="cp1252"),
]
TAL = [
    dict(ta="Oncology", region="GLOBAL", year=2024, fmt="pptx", image_slide=True),
    dict(ta="Cardiology", region="EU", year=2023, fmt="pptx"),
    dict(ta="Neurology", region="US", year=2024, fmt="pptx", unordered_boxes=True),
    dict(ta="Immunology", region="GLOBAL", year=2023, fmt="pptx"),
    dict(ta="Immunology", region="EU", year=2022, fmt="pptx"),
]

DEFAULT_SENS = {"HCP Perception Survey": "INTERNAL", "Competitive Landscape Report": "RESTRICTED",
                "Indication Deep Dive": "PUBLIC", "Market Access Insight": "INTERNAL",
                "KOL Interview Summary": "RESTRICTED", "Therapy Area Landscape": "INTERNAL"}


def _title(g, st, spec):
    if st == "HCP Perception Survey":
        return f"HCP Perception Survey: {spec['brands'][0]} in {spec['indication']} ({spec['region']}, {spec['year']})"
    if st == "Competitive Landscape Report":
        return f"Competitive Landscape: {spec['indication']} ({spec['region']}, {spec['year']})"
    if st == "Indication Deep Dive":
        return f"Indication Deep Dive: {spec['indication']} ({spec['region']}, {spec['year']})"
    if st == "Market Access Insight":
        return f"Market Access Insight: {spec['brands'][0]} in {cname(spec['country'])} ({spec['year']})"
    if st == "KOL Interview Summary":
        return f"KOL Interview Summary: {spec['brands'][0]} in {cname(spec['country'])} ({spec['year']})"
    return f"Therapy Area Landscape: {spec['ta']} ({spec['region']}, {spec['year']})"


def make_specs(rng):
    g = Gen(rng)
    specs = []

    def add(st, p):
        spec = dict(p)
        spec["study_type"] = st
        if "brand" in spec:
            spec["brands"] = [spec.pop("brand")]
        specs.append(spec)

    for p in HCP:
        add("HCP Perception Survey", p)
    for p in CL:
        add("Competitive Landscape Report", p)
    for p in ID:
        add("Indication Deep Dive", p)
    for p in MA:
        add("Market Access Insight", p)
    for p in KOL:
        add("KOL Interview Summary", p)
    for p in TAL:
        add("Therapy Area Landscape", p)
    return g, specs


def finish_spec(g, spec, st, idx):
    """Fill derived fields: richness, indication, author, date, region."""
    r = g.rng
    spec["richness"] = spec.pop("rich", None) or g.pick(["low", "med", "high"])
    if "ind" in spec:
        spec["indication"] = spec.pop("ind")
    if st in ("HCP Perception Survey", "Market Access Insight", "KOL Interview Summary"):
        b = vocab.BRANDS[spec["brands"][0]]
        spec.setdefault("indication", b["indications"][0])
        spec["therapeutic_area"] = b["ta"]
    if st == "Market Access Insight":
        spec["region"] = vocab.GEOGRAPHIES[spec["country"]]["region"]
    if st == "KOL Interview Summary":
        spec["region"] = vocab.GEOGRAPHIES[spec["country"]]["region"]
    if st == "Therapy Area Landscape":
        spec["therapeutic_area"] = spec["ta"]
    if spec.get("no_author"):
        spec["author"] = None
    else:
        spec["author"] = vocab.ANALYST_NAMES[idx % len(vocab.ANALYST_NAMES)]
    spec["report_date"] = (date(spec["year"], r.randint(1, 12), r.randint(1, 28))).isoformat()
    spec["sensitivity"] = spec.pop("sens", None) or DEFAULT_SENS[st]
    spec["n"] = spec.get("n", 0)


def render_any(spec, st, fmt, sections, table, notes, path):
    if fmt == "pdf":
        (render_scanned_pdf if spec.get("scanned") else render_pdf)(spec, sections, table, path)
    elif fmt == "docx":
        render_docx(spec, sections, table, path)
    elif fmt == "pptx":
        render_pptx(spec, sections, table, notes, path)
    else:
        render_txt(spec, sections, table, path)


def generate(out: Path, seed: int = SEED) -> list:
    rng = random.Random(seed)
    g, specs = make_specs(rng)
    out = Path(out)
    if out.exists():
        for p in out.iterdir():
            if p.is_file():
                p.unlink()
    out.mkdir(parents=True, exist_ok=True)
    manifest, built = [], {}

    for idx, spec in enumerate(specs, start=1):
        st = spec["study_type"]
        finish_spec(g, spec, st, idx)
        sections, table, notes = BUILDERS[st](g, spec)
        spec["title"] = _title(g, st, spec)
        fmt = spec["fmt"]
        sections = pad(g, spec, sections, fmt)
        name = f"study_{idx:03d}.{fmt}"
        render_any(spec, st, fmt, sections, table, notes, out / name)
        spec["filename"] = name
        built[idx] = (copy.deepcopy(spec), sections, table, notes)
        manifest.append(spec)

    nxt = len(specs)
    # near-duplicate: same title/author/date, a few words changed
    for idx, (spec, sections, table, notes) in list(built.items()):
        if spec.get("near_dup_source"):
            nxt += 1
            s2 = copy.deepcopy(spec)
            s2.pop("near_dup_source", None)
            sec2 = [(h, [p.replace("moderately", "reasonably").replace("main", "primary") for p in ps]) for h, ps in sections]
            sec2[0][1][0] = sec2[0][1][0].replace("This survey", "This survey")
            name = f"study_{idx:03d}_rev.{spec['fmt']}"
            render_any(s2, s2["study_type"], s2["fmt"], sec2, table, notes, out / name)
            s2.update(filename=name, planted={"near_duplicate_of": spec["filename"]}, ingest_age=0)
            manifest.append(s2)
    # new version: same title and author, later report date, different content, other format
    for idx, (spec, sections, table, notes) in list(built.items()):
        if spec.get("new_version_source"):
            s2 = copy.deepcopy(spec)
            s2.pop("new_version_source", None)
            s2["report_date"] = (date.fromisoformat(spec["report_date"]) + timedelta(days=420)).isoformat()
            s2["year"] = spec["year"] + 1
            s2["fmt"], s2["image_slide"] = "pdf", False
            g2 = Gen(random.Random(seed + 99))
            sec2, table2, notes2 = build_cl(g2, s2)
            sec2 = pad(g2, s2, sec2, "pdf")
            s2["title"] = spec["title"]                       # same title => version candidate
            name = f"study_{idx:03d}_update.pdf"
            render_any(s2, s2["study_type"], "pdf", sec2, table2, notes2, out / name)
            s2.update(filename=name, planted={"new_version_of": spec["filename"]}, ingest_age=0)
            manifest.append(s2)
    # exact duplicate: byte-identical copy
    for idx, (spec, *_rest) in list(built.items()):
        if spec.get("exact_dup_source"):
            name = f"study_{idx:03d}_copy.{spec['fmt']}"
            shutil.copyfile(out / spec["filename"], out / name)
            s2 = copy.deepcopy(spec)
            s2.pop("exact_dup_source", None)
            s2.update(filename=name, planted={"exact_duplicate_of": spec["filename"]}, ingest_age=0)
            manifest.append(s2)

    for spec in manifest:
        p = spec.setdefault("planted", {})
        if spec.get("scanned"):
            p["scanned_pdf"] = True
        if spec.get("pii"):
            p["pii"] = True
        spec.setdefault("ingest_age", spec.pop("age", None))
        spec["ingest_age"] = spec.get("ingest_age")
        spec.pop("age", None)
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return manifest


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(config.DATA_DIR))
    ap.add_argument("--seed", type=int, default=SEED)
    a = ap.parse_args()
    m = generate(Path(a.out), a.seed)
    files = [x["filename"] for x in m]
    by_type = {}
    for x in m:
        by_type.setdefault(x["study_type"], 0)
        by_type[x["study_type"]] += 1
    by_ext = {}
    for fn in files:
        by_ext[Path(fn).suffix] = by_ext.get(Path(fn).suffix, 0) + 1
    print(f"Generated {len(files)} files in {a.out}")
    print("by study type:", by_type)
    print("by format:", by_ext)


if __name__ == "__main__":
    main()


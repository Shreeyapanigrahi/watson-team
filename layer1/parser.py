"""Multi-format document parser (Section 4.1).

Every format yields the same ParsedDocument, so nothing downstream sees a raw
file. `units` is an addition to the brief's field list: the structural pieces of
the document (page / heading section / slide / paragraph group) with character
offsets into raw_text, which the structural chunking strategy needs.
"""
import math
import re
from dataclasses import dataclass, field, asdict
from datetime import datetime
from pathlib import Path

import chardet
import pymupdf
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation

from .errors import UnsupportedFormatError

SUPPORTED = {".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".txt": "txt"}


@dataclass
class ParsedDocument:
    raw_text: str
    page_or_slide_count: int          # pages (pdf), slides (pptx), est. pages (docx), lines (txt)
    format: str
    source_filename: str
    extracted_metadata: dict = field(default_factory=dict)
    parse_warnings: list = field(default_factory=list)
    units: list = field(default_factory=list)

    def to_dict(self) -> dict:
        return asdict(self)

    @property
    def word_count(self) -> int:
        return len(self.raw_text.split())


class _Builder:
    """Joins unit texts with a blank line and records char offsets."""

    def __init__(self):
        self.parts, self.units, self.pos = [], [], 0

    def add(self, text: str, **meta):
        text = text.strip()
        if not text:
            return
        if self.parts:
            self.pos += 2          # the "\n\n" separator
        start = self.pos
        self.parts.append(text)
        self.pos += len(text)
        self.units.append({**meta, "start": start, "end": self.pos})

    @property
    def text(self) -> str:
        return "\n\n".join(self.parts)


_PLACEHOLDER_META = {"", "untitled", "(untitled)", "anonymous", "(anonymous)", "unknown"}


def _clean(value) -> str | None:
    """Writers such as reportlab fill empty properties with 'untitled' / '(anonymous)'."""
    value = (value or "").strip()
    return None if value.lower() in _PLACEHOLDER_META else value


def _iso_date(value) -> str | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", str(value))
    if m:
        return f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
    m = re.match(r"(\d{4}-\d{2}-\d{2})", str(value))
    return m.group(1) if m else None


# ---------------------------------------------------------------- PDF
def _pdf_page_text(page) -> tuple[str, bool]:
    """Page text in reading order. Detects a two-column layout and reads the
    left column before the right one (plain extraction can interleave them)."""
    width = page.rect.width
    blocks = [b for b in page.get_text("blocks") if b[6] == 0 and b[4].strip()]
    left = [b for b in blocks if b[2] < width * 0.55]
    right = [b for b in blocks if b[0] > width * 0.45]
    two_col = False
    if len(left) >= 2 and len(right) >= 2:
        two_col = any(l[1] < r[3] and r[1] < l[3] for l in left for r in right)
    if two_col:
        ordered = sorted(left, key=lambda b: b[1]) + sorted(right, key=lambda b: b[1])
        rest = [b for b in blocks if b not in left and b not in right]
        ordered = sorted(rest, key=lambda b: b[1])[:1] + ordered if rest else ordered
        return "\n".join(b[4].strip() for b in ordered), True
    return page.get_text("text", sort=True).strip(), False


def _parse_pdf(path: Path) -> ParsedDocument:
    doc = pymupdf.open(path)
    warnings, builder = [], _Builder()
    total_chars = 0
    for i, page in enumerate(doc, start=1):
        text, two_col = _pdf_page_text(page)
        total_chars += len(text)
        if two_col:
            warnings.append(f"MULTI_COLUMN: page {i} has a multi-column layout; read column by column, verify order")
        if len(text) < 10:
            warnings.append(f"EMPTY_PAGE: page {i} has no extractable text")
        builder.add(text, kind="page", title=f"Page {i}", level=1)
    n_pages = len(doc)
    if n_pages and total_chars < 20 * n_pages:
        warnings.insert(0, f"SCANNED_PDF: no text layer found on {n_pages} page(s); "
                           "likely scanned images. OCR is not performed, so no content is extracted")
    md = doc.metadata or {}
    meta = {
        "title": _clean(md.get("title")),
        "author": _clean(md.get("author")),
        "created": _iso_date(md.get("creationDate")),
    }
    doc.close()
    return ParsedDocument(builder.text, n_pages, "pdf", path.name, meta, warnings, builder.units)


# ---------------------------------------------------------------- DOCX
def _para_text(p) -> str:
    # Reads every w:t (including inserted runs) and skips w:delText, so tracked-change
    # deletions are stripped before extraction.
    return "".join(t.text or "" for t in p._p.iter(qn("w:t")))


def _iter_blocks(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _parse_docx(path: Path) -> ParsedDocument:
    doc = DocxDocument(path)
    warnings = []
    if doc.element.body.findall(".//" + qn("w:del")) or doc.element.body.findall(".//" + qn("w:ins")):
        warnings.append("TRACKED_CHANGES: tracked-change markup stripped; deleted text excluded")

    sections, headings = [], []
    current = {"title": None, "level": 0, "lines": []}

    def flush():
        if current["lines"]:
            sections.append((current["title"], current["level"], "\n".join(current["lines"])))

    for block in _iter_blocks(doc):
        if isinstance(block, Paragraph):
            text = _para_text(block).strip()
            if not text:
                continue
            style = (block.style.name or "") if block.style is not None else ""
            m = re.match(r"Heading (\d)", style)
            if m:
                level = int(m.group(1))
                headings.append({"level": level, "text": text})
                if level <= 2:              # H1/H2 start a new section; H3 stays inside it
                    flush()
                    current = {"title": text, "level": level, "lines": [text]}
                    continue
            elif style == "Title":
                headings.append({"level": 0, "text": text})
            current["lines"].append(text)
        else:                               # table: iterate cells, never merge into paragraphs
            seen = set()
            for row in block.rows:
                cells = []
                for cell in row.cells:
                    if cell._tc in seen:         # merged cells repeat; hold the element, not id(), which is reused
                        continue
                    seen.add(cell._tc)
                    cells.append(" ".join(_para_text(p).strip() for p in cell.paragraphs).strip())
                line = " | ".join(c for c in cells if c)
                if line:
                    current["lines"].append(line)
    flush()

    builder = _Builder()
    for title, level, text in sections:
        builder.add(text, kind="section", title=title or "Introduction", level=level)

    cp = doc.core_properties
    meta = {
        "title": _clean(cp.title),
        "author": _clean(cp.author),
        "created": _iso_date(cp.created),
        "headings": headings,
        "page_count_estimated": True,
    }
    pages = max(1, math.ceil(len(builder.text.split()) / 500))
    return ParsedDocument(builder.text, pages, "docx", path.name, meta, warnings, builder.units)


# ---------------------------------------------------------------- PPTX
def _flatten(shapes):
    for sh in shapes:
        if sh.shape_type == 6:              # group
            yield from _flatten(sh.shapes)
        else:
            yield sh


def _parse_pptx(path: Path) -> ParsedDocument:
    prs = Presentation(path)
    warnings, builder = [], _Builder()
    for idx, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        frames = []
        for sh in _flatten(slide.shapes):
            if sh is title_shape:
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                frames.append((sh.top or 0, sh.left or 0, sh.text_frame.text.strip()))
            elif getattr(sh, "has_table", False) and sh.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells if c.text.strip()) for r in sh.table.rows]
                frames.append((sh.top or 0, sh.left or 0, "\n".join(r for r in rows if r)))
        frames.sort(key=lambda t: (t[0], t[1]))         # text boxes are unordered: top-left first
        body = "\n".join(f[2] for f in frames)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if not title and not body and not notes:
            warnings.append(f"NON_TEXTUAL_SLIDE: slide {idx} has no extractable text (chart or image only)")
            continue
        parts = [f"Slide {idx}: {title}".strip(": ").strip()] if title else [f"Slide {idx}"]
        if body:
            parts.append(body)
        if notes:
            parts.append(f"Speaker notes: {notes}")
        builder.add("\n".join(parts), kind="slide", title=title or f"Slide {idx}", level=1, slide=idx)
    cp = prs.core_properties
    first_title = next((u["title"] for u in builder.units if u["title"] and not u["title"].startswith("Slide ")), None)
    meta = {
        "title": _clean(cp.title) or first_title,
        "author": _clean(cp.author),
        "created": _iso_date(cp.created),
    }
    return ParsedDocument(builder.text, len(prs.slides), "pptx", path.name, meta, warnings, builder.units)


# ---------------------------------------------------------------- TXT
def _parse_txt(path: Path) -> ParsedDocument:
    raw = path.read_bytes()
    det = chardet.detect(raw)
    enc, conf = det.get("encoding") or "utf-8", det.get("confidence") or 0.0
    warnings = []
    try:
        text = raw.decode(enc)
    except (UnicodeDecodeError, LookupError):
        text = raw.decode("utf-8", errors="replace")
        warnings.append(f"ENCODING_FALLBACK: detected '{enc}' failed; decoded as UTF-8 with replacement")
    if conf < 0.6:
        warnings.append(f"LOW_ENCODING_CONFIDENCE: '{enc}' detected with confidence {conf:.2f}")
    text = text.lstrip("﻿").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text).strip()          # normalise blank lines
    builder = _Builder()
    for para in re.split(r"\n\s*\n", text):
        builder.add(para, kind="paragraph", title=None, level=1)
    meta = {
        "title": None,                    # TXT has no properties; the metadata extractor guesses from the first line
        "author": None,
        "created": None,
        "encoding": enc,
        "encoding_confidence": round(conf, 2),
    }
    return ParsedDocument(builder.text, len(text.splitlines()), "txt", path.name, meta, warnings, builder.units)


_PARSERS = {"pdf": _parse_pdf, "docx": _parse_docx, "pptx": _parse_pptx, "txt": _parse_txt}


def parse_file(path) -> ParsedDocument:
    path = Path(path)
    fmt = SUPPORTED.get(path.suffix.lower())
    if fmt is None:
        raise UnsupportedFormatError(
            f"Unsupported file type '{path.suffix}'. Supported: {', '.join(sorted(SUPPORTED))}")
    return _PARSERS[fmt](path)


if __name__ == "__main__":
    import json
    import sys
    pd = parse_file(sys.argv[1])
    print(f"format={pd.format} words={pd.word_count} pages/slides/lines={pd.page_or_slide_count} "
          f"units={len(pd.units)}")
    print("warnings:", json.dumps(pd.parse_warnings, indent=2))
    print("metadata:", {k: v for k, v in pd.extracted_metadata.items() if k != "headings"})
